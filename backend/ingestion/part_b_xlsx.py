"""Part B spreadsheet ingestion.
"""

from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict
from app.config import get_settings
from app.logging import get_logger
from .chunking import Chunk, ChunkKind

log = get_logger(__name__)

HEADER_ROWS = 6
ROWS_PER_CHUNK = 45
MIN_SHEET_CHARS = 40

STATEMENT_KINDS = {
    "BS": "balance_sheet",
    "PL": "income_statement",
    "EQ": "changes_in_equity",
    "CF": "cash_flow",
}


class SheetRecord(BaseModel):
    model_config = ConfigDict(frozen=True)

    sheet: str
    statement: str
    rows: int
    chars: int
    text_extractable: bool
    reason: str | None = None


def statement_kind(sheet_name: str) -> str:
    """`BS-2-4` -> balance_sheet, `PL 5` -> income_statement."""
    prefix = re.split(r"[\s\-]", sheet_name.strip())[0].upper()
    return STATEMENT_KINDS.get(prefix, "other")


def _column_letter(index: int) -> str:
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters or "A"


def _row_text(row: tuple) -> str:
    return " | ".join("" if v is None else str(v).strip() for v in row).strip(" |")


def read_sheets(path: Path) -> list[tuple[str, list[tuple]]]:
    """Every sheet as a list of row tuples, values only."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        out: list[tuple[str, list[tuple]]] = []
        for name in workbook.sheetnames:
            worksheet = workbook[name]
            rows = [r for r in worksheet.iter_rows(values_only=True)
                    if any(c is not None and str(c).strip() for c in r)]
            out.append((name, rows))
        return out
    finally:
        workbook.close()


def workbook_text(path: Path, limit_rows: int = 40) -> str:
    """The front of every sheet, for entity, currency and period detection."""
    parts: list[str] = []
    for name, rows in read_sheets(path):
        parts.append(name)
        parts.extend(_row_text(r) for r in rows[:limit_rows])
    return "\n".join(parts)


def chunk_workbook(
    path: Path, *, company_code: str, filename: str, doc_type: str
) -> tuple[list[SheetRecord], list[Chunk]]:
    records: list[SheetRecord] = []
    chunks: list[Chunk] = []

    for sheet_name, rows in read_sheets(path):
        text_rows = [_row_text(r) for r in rows]
        total_chars = sum(len(t) for t in text_rows)
        usable = total_chars >= MIN_SHEET_CHARS and bool(rows)
        records.append(
            SheetRecord(
                sheet=sheet_name,
                statement=statement_kind(sheet_name),
                rows=len(rows),
                chars=total_chars,
                text_extractable=usable,
                reason=None if usable else "sheet_empty",
            )
        )
        if not usable:
            continue

        width = max((len(r) for r in rows), default=1)
        last_column = _column_letter(width)
        header = text_rows[:HEADER_ROWS]
        heading = f"{filename} · {sheet_name} ({statement_kind(sheet_name)})"

        ordinal = 0
        for start in range(HEADER_ROWS, len(rows), ROWS_PER_CHUNK):
            block = text_rows[start : start + ROWS_PER_CHUNK]
            if not any(b.strip() for b in block):
                continue
            ordinal += 1
            first_row, last_row = start + 1, min(start + ROWS_PER_CHUNK, len(rows))
            cell_range = f"A{first_row}:{last_column}{last_row}"
            body = "\n".join([*header, "", *block]).strip()

            chunks.append(
                Chunk(
                    chunk_id=f"{company_code}|{filename}|{sheet_name}|c{ordinal:02d}",
                    company_code=company_code,
                    filename=filename,
                    doc_type=doc_type,
                    page=0,  # not a paginated document; the sheet is the locator
                    ordinal=ordinal,
                    kind=ChunkKind.TABLE,
                    heading=heading,
                    text=body,
                    chars=len(body),
                    sheet=sheet_name,
                    cell_range=cell_range,
                )
            )

    log.info(
        "xlsx.chunked",
        file=filename,
        sheets=len(records),
        usable=sum(1 for r in records if r.text_extractable),
        chunks=len(chunks),
    )
    return records, chunks


def workbook_paths(part_b_dir: Path | None = None) -> list[Path]:
    """Workbooks under Part B, deduplicated.
    """
    part_b_dir = part_b_dir or get_settings().part_b_dir
    return sorted({p.resolve() for pattern in ("*.xlsx", "*.XLSX")
                   for p in part_b_dir.rglob(pattern)})
